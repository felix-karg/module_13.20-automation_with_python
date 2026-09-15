import openpyxl

# Open file and save it in a variable
inv_file = openpyxl.load_workbook("inventory.xlsx")
# Save content of the sheet named 'Sheet1' in another variable
product_list = inv_file["Sheet1"]


### Exercise1: Calculate number of products per supplier ###
products_per_supplier = {}
# Start loop from second row to skip headers
for product_row in range(2, product_list.max_row + 1):
    # extract value of 4th column of each row
    supplier_name = product_list.cell(product_row, 4).value

    if supplier_name in products_per_supplier:
        products_per_supplier[supplier_name] = products_per_supplier[supplier_name] + 1
    else:
        products_per_supplier[supplier_name] = 1

print(products_per_supplier)